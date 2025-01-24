from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


def add_rectangle(slide, left, top, width, height, text, color, font_size=12):
    """
    添加矩形框并设置文本和颜色
    """
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*color)  # 填充颜色
    shape.line.color.rgb = RGBColor(0, 0, 0)      # 边框颜色
    text_frame = shape.text_frame
    text_frame.text = text
    for paragraph in text_frame.paragraphs:
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(0, 0, 0)  # 黑色字体
    return shape


def add_arrow(slide, start_left, start_top, end_left, end_top):
    """
    添加箭头以连接形状
    """
    connector = slide.shapes.add_connector(
        MSO_SHAPE.RIGHT_ARROW, start_left, start_top, end_left, end_top
    )
    connector.line.color.rgb = RGBColor(0, 0, 0)  # 黑色箭头
    connector.line.width = Pt(2)  # 箭头宽度
    connector.line.end_arrowhead = True


def create_framework_ppt():
    """
    创建框架图的 PowerPoint 幻灯片
    """
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 添加空白幻灯片

    rect_width = Inches(1.2)
    rect_height = Inches(0.6)

    # 绘制 view 1 和 view 2
    view_box = add_rectangle(slide, Inches(1), Inches(1), Inches(2), Inches(1.2), "View 1\nView 2", color=(255, 255, 255), font_size=12)
    # add_arrow(slide, Inches(3), Inches(1.6), Inches(4), Inches(1.6))

    clip_box = add_rectangle(slide, Inches(4), Inches(1.2), rect_width, rect_height, "clip", color=(135, 206, 250), font_size=12)

    # VAE 部分
    vae_box = add_rectangle(slide, Inches(1), Inches(3), rect_width / 2, rect_height, "VAE", color=(135, 206, 250), font_size=12)
    # add_arrow(slide, Inches(2), Inches(3.3), Inches(2.8), Inches(3.3))

    repeat_box = add_rectangle(slide, Inches(3), Inches(3), rect_width, rect_height, "Repeat n times", color=(135, 206, 250), font_size=12)

    # noise latents 部分
    latents_box = add_rectangle(slide, Inches(4), Inches(4.5), rect_width, rect_height, "Noise Latents", color=(135, 206, 250), font_size=12)
    c_circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.5), Inches(4.7), Inches(0.6), Inches(0.6))
    c_circle.fill.solid()
    c_circle.fill.fore_color.rgb = RGBColor(135, 206, 250)
    c_circle.text = "c"

    # add_arrow(slide, Inches(3.5), Inches(3.3), Inches(3.5), Inches(4.7))
    # add_arrow(slide, Inches(5.05), Inches(5), Inches(6), Inches(5))

    # embedding 部分
    embedding1_box = add_rectangle(slide, Inches(5.9), Inches(2), rect_width, rect_height, "Embedding 1", color=(144, 238, 144), font_size=12)
    embedding2_box = add_rectangle(slide, Inches(5.9), Inches(3), rect_width, rect_height, "Embedding 2", color=(144, 238, 144), font_size=12)

    # conv blocks
    conv_blocks_box = add_rectangle(slide, Inches(7.5), Inches(3.5), rect_width, rect_height, "Conv Blocks", color=(135, 206, 250), font_size=12)
    spatial_blocks = add_rectangle(slide, Inches(9), Inches(4), rect_width, rect_height, "Spatial Atten", color=(255, 182, 193), font_size=12)

    temporal_blocks = add_rectangle(slide, Inches(11), Inches(4), rect_width, rect_height, "Temporal Atten", color=(255, 255, 0), font_size=12)
    view_blocks = add_rectangle(slide, Inches(13), Inches(4), rect_width, rect_height, "View Atten", color=(255, 255, 0), font_size=12)

    # 添加箭头连接各部分
    # add_arrow(slide, Inches(6.5), Inches(2.4), Inches(7.2), Inches(3.8))
    # add_arrow(slide, Inches(8), Inches(4.3), Inches(9.3), Inches(4.3))
    # add_arrow(slide, Inches(10.8), Inches(4.3), Inches(11.8), Inches(4.3))
    # add_arrow(slide, Inches(12.8), Inches(4.3), Inches(14), Inches(4.3))

    # 保存PPT
    prs.save('framework_diagram.pptx')
    print("框架图已保存为 'framework_diagram.pptx'")


# 运行生成框架图
create_framework_ppt()