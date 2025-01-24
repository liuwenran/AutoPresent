from pptx import Presentation

# 创建一个新的 PowerPoint 对象
presentation = Presentation()

# 添加一个标题幻灯片
slide = presentation.slides.add_slide(presentation.slide_layouts[0])

# 设置标题和副标题（标注为 placeholders）
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Hello, Python-pptx"
subtitle.text = "This is a simple example of python-pptx."

# 保存文件
presentation.save("example.pptx")